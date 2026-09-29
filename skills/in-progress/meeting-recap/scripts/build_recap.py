#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1,<2"]
# ///
"""Build a meeting recap's PDF from its markdown, in the brand set in brand/brand.md.

    uv run scripts/build_recap.py RECAP.md [--brand FOLDER]

Reads only the recap shape. Writes RECAP.docx beside RECAP.md, converts it to RECAP.pdf with
LibreOffice and removes the .docx. The page is described in references/layout.md. Any problem
stops the build with one line naming it. No network access; writes only beside the recap."""

from __future__ import annotations

import argparse
import colorsys
from dataclasses import dataclass, field
from pathlib import Path
import re
import shutil
import subprocess
import sys

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.image.image import Image
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Mm, Pt, RGBColor
except ImportError:
    sys.exit("build_recap: python-docx is not installed. Run `uv run scripts/build_recap.py "
             "RECAP.md`, or install it with `python3 -m pip install python-docx` and run again.")

SKILL_DIR = Path(__file__).resolve().parent.parent
AI_NOTE = "A Note On This Summary"
DEFAULTS = {
    "configured": "no", "organization": "", "tagline": "", "contact_line": "",
    "document_label": "MEETING RECAP", "display_font": "Arial", "body_font": "Arial",
    "accent_color": "#DBE5F1", "text_color": "#222222", "logo": "", "page_size": "Letter",
    "footer": "{organization} · {contact_line}", "style_notes": "",
}
FOOTER_FIELDS = ("organization", "tagline", "contact_line")
PAGE_SIZES = {"letter": (Inches(8.5), Inches(11)), "a4": (Mm(210), Mm(297))}
GRAY, HAIRLINE = "666666", "C8C8C8"
ACTION_WIDTHS = (0.06, 0.35, 0.12, 0.19, 0.28)  # #, Item, Owner, Target or Completed, Notes
LOGO_HEIGHT, LOGO_MAX_WIDTH = Inches(0.6), Inches(2.5)
KNOWN_SOFFICE = ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe")

# Word rejects properties out of schema order, so each insertion names what may follow it.
AFTER_RUN_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u")
AFTER_PARA_BORDER = ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr")
AFTER_TABLE_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
AFTER_CELL_SHADE = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign")


class Fail(Exception):
    """A problem the user can fix, reported as one line."""


# ---------------------------------------------------------------- brand card
def load_brand(folder: Path) -> dict:
    card = folder / "brand.md"
    if not card.is_file():
        raise Fail(f"no brand card at {card}")
    brand = dict(DEFAULTS)
    for line in card.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^([a-z_]+):[ \t]*(.*?)[ \t]*$", line)
        if match and match.group(1) in DEFAULTS:
            brand[match.group(1)] = match.group(2)
        elif match:
            print(f"build_recap: note: brand.md field '{match.group(1)}' is not one this script "
                  "reads; ignored.", file=sys.stderr)
    for key in ("document_label", "display_font", "body_font", "page_size", "accent_color",
                "text_color"):
        brand[key] = brand[key] or DEFAULTS[key]
    for key in ("accent_color", "text_color"):
        if not re.fullmatch(r"#?[0-9A-Fa-f]{6}", brand[key]):
            raise Fail(f"brand.md {key} must be a hex color such as #1F3A5F, not '{brand[key]}'")
        brand[key] = brand[key].lstrip("#").upper()
    if brand["page_size"].lower() not in PAGE_SIZES:
        raise Fail(f"brand.md page_size must be Letter or A4, not '{brand['page_size']}'")
    brand["logo_path"] = folder / brand["logo"] if brand["logo"] else None
    if brand["logo_path"] and not brand["logo_path"].is_file():
        raise Fail(f"brand.md names the logo '{brand['logo']}', but {brand['logo_path']} is not there")
    if brand["logo_path"] and brand["logo_path"].suffix.lower() not in (".png", ".jpg", ".jpeg"):
        raise Fail(f"the logo must be a PNG or JPG file, not '{brand['logo']}'")
    brand["footer_text"] = footer_text(brand)
    return brand


def footer_text(brand: dict) -> str:
    """Fill the footer template, dropping empty fields and the separators beside them."""
    def fill(match: re.Match) -> str:
        if match.group(1) not in FOOTER_FIELDS:
            raise Fail(f"brand.md footer uses {{{match.group(1)}}}, which is not a field it can show")
        return brand[match.group(1)]

    text = re.sub(r"\{([a-z_]+)\}", fill, brand["footer"])
    separator = next(iter(re.findall(r"[·|]", text)), "·")
    return f" {separator} ".join(part.strip() for part in re.split(r"[·|]", text) if part.strip())


def rgb(hex6: str) -> tuple[float, ...]:
    return tuple(int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))


def luminance(hex6: str) -> float:
    red, green, blue = (c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
                        for c in rgb(hex6))
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def contrast(one: str, two: str) -> float:
    light, dark = sorted((luminance(one), luminance(two)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def readable(hex6: str, ground: str = "FFFFFF", ratio: float = 4.5) -> str:
    """The color, darkened in its own hue and saturation until it reads on the ground color."""
    hue, lightness, saturation = colorsys.rgb_to_hls(*rgb(hex6))
    while contrast(hex6, ground) < ratio and lightness > 0.02:
        lightness -= 0.02
        hex6 = "".join(f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(hue, lightness, saturation))
    return hex6


def pale(hex6: str) -> str:
    """The color, mixed toward white until it is a pale fill."""
    for share in (step / 20 for step in range(21)):
        mixed = "".join(f"{round(c * 255 * (1 - share) + 255 * share):02X}" for c in rgb(hex6))
        if luminance(mixed) >= 0.8:
            return mixed
    return "FFFFFF"


# ---------------------------------------------------------------- the recap shape
@dataclass
class Block:
    kind: str  # para, bullet, number or table
    text: str = ""
    number: str = ""
    rows: list[list[str]] = field(default_factory=list)


FACT = re.compile(r"^\*\*(?P<key>[^*]+?)(?::\*\*|\*\*:)\s*(?P<value>.+)$")
BULLET = re.compile(r"^[-*+]\s+(.+)$")
NUMBER = re.compile(r"^(\d{1,3})[.)]\s+(.+)$")
SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
SHAPE_HINT = "use ## sections, paragraphs, - bullets, 1. items and | tables"


def parse(path: Path) -> tuple[str, list[tuple[str, str]], list[tuple[str, list[Block]]]]:
    """The title, the facts and the sections, or a Fail naming the first line out of shape."""
    lines = path.read_text(encoding="utf-8").splitlines()
    title, facts, sections = None, [], []
    blocks: list[Block] | None = None
    current: Block | None = None

    def refuse(number: int, why: str) -> None:
        raise Fail(f"{path.name} line {number} is not part of the recap shape ({why}): "
                   f"'{lines[number - 1].strip()[:70]}'")

    index = 0
    while index < len(lines):
        number, raw = index + 1, lines[index]
        text = raw.strip()
        index += 1
        if not text:
            current = None
            continue
        indented = len(raw) - len(raw.lstrip(" \t")) >= 2
        if text.startswith("# ") and title is None:
            title = text[2:].strip()
            continue
        if title is None:
            refuse(number, "the recap opens with its # title")
        if text.startswith("## "):
            blocks, current = [], None
            sections.append((text[3:].strip(), blocks))
            continue
        if text.startswith("#"):
            refuse(number, SHAPE_HINT)
        if blocks is None:
            fact = FACT.match(text)
            if not fact:
                refuse(number, "only **Key:** value lines sit between the title and the first section")
            facts.append((fact["key"].strip(), fact["value"].strip()))
            continue
        if text.startswith("|"):
            rows = [(number, text)]
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append((index + 1, lines[index].strip()))
                index += 1
            blocks.append(parse_table(rows, refuse))
            current = None
            continue
        if text.startswith((">", "```", "~~~", "<")) or re.fullmatch(r"[-*_ ]{3,}", text):
            refuse(number, SHAPE_HINT)
        bullet, numbered = BULLET.match(text), NUMBER.match(text)
        if (bullet or numbered) and indented and current and current.kind != "para":
            refuse(number, "lists are one level deep")
        if bullet:
            current = Block("bullet", bullet.group(1))
        elif numbered:
            current = Block("number", numbered.group(2), number=numbered.group(1))
        elif current:  # a wrapped line continues the paragraph or item above it
            current.text += " " + text
            continue
        else:
            current = Block("para", text)
        blocks.append(current)

    if title is None:
        raise Fail(f"{path.name} has no # title line")
    if not sections:
        raise Fail(f"{path.name} has no ## sections")
    if sections[-1][0].lower() != AI_NOTE.lower():
        raise Fail(f"{path.name} must end with the '## {AI_NOTE}' section; the last section is "
                   f"'## {sections[-1][0]}'")
    return title, facts, sections


def split_row(text: str) -> list[str]:
    text = text.strip()
    text = text[1:] if text.startswith("|") else text
    text = text[:-1] if text.endswith("|") and not text.endswith("\\|") else text
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", text)]


def parse_table(rows: list[tuple[int, str]], refuse) -> Block:
    if len(rows) < 2 or not SEPARATOR.match(rows[1][1]):
        refuse(rows[0][0], "a table needs a header row, then a |---| line")
    header, body = split_row(rows[0][1]), []
    for number, text in rows[2:]:
        cells = split_row(text)
        if len(cells) != len(header):
            refuse(number, f"this row has {len(cells)} cells and the header has {len(header)}")
        body.append(cells)
    return Block("table", rows=[header] + body)


# ---------------------------------------------------------------- text
INLINE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|\[(?P<label>[^\]]+)\]\((?P<url>[^)\s]+)\)"
    r"|(?<![\w*])\*(?P<em>[^*\s](?:[^*]*[^*\s])?)\*(?![\w*])"
    r"|`(?P<code>[^`]+)`"
)


def unescape(text: str) -> str:
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|])", r"\1", text)


def plain(text: str) -> str:
    """The words a cell shows, without markdown, for measuring."""
    return unescape(re.sub(r"\[([^\]]+)\]\([^)]*\)|\*\*|`", lambda m: m.group(1) or "", text))


def make(tag: str, **attributes: str):
    """A WordprocessingML element with w: attributes."""
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), value)
    return node


def style(run, font: str, size: float, color: str, bold=False, italic=False, underline=False,
          spacing: int = 0, link: str | None = None) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().attrib.update(
        {qn("w:eastAsia"): font, qn("w:cs"): font})
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold, run.font.italic, run.font.underline = bold, italic, underline
    if spacing:  # letter-spacing, in twentieths of a point
        tracking = make("w:spacing", val=str(spacing))
        run._element.get_or_add_rPr().insert_element_before(tracking, *AFTER_RUN_SPACING)


def add_text(paragraph, text: str, look: dict, bold=False) -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            style(paragraph.add_run(unescape(text[position:match.start()])), **look, bold=bold)
        if match.group("bold") is not None:
            add_text(paragraph, match.group("bold"), look, bold=True)
        elif match.group("url") is not None:
            add_link(paragraph, match.group("url"), unescape(match.group("label")), look, bold)
        elif match.group("code") is not None:
            style(paragraph.add_run(match.group("code")), **look, bold=bold)
        else:
            style(paragraph.add_run(unescape(match.group("em"))), **look, bold=bold, italic=True)
        position = match.end()
    if position < len(text):
        style(paragraph.add_run(unescape(text[position:])), **look, bold=bold)


def add_link(paragraph, url: str, label: str, look: dict, bold: bool) -> None:
    """Web and mail links stay live; a local file link keeps its text only."""
    if not re.match(r"^(https?://|mailto:)", url, re.I):
        style(paragraph.add_run(label), **look, bold=bold)
        return
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = paragraph.add_run(label)
    style(run, look["font"], look["size"], look["link"], bold=bold, underline=True)
    link.append(run._element)
    paragraph._p.append(link)


def spacing(paragraph, before=0.0, after=6.0, line=1.15, keep=False):
    """Paragraph spacing in points; a line given in points is exact, which makes a spacer."""
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), line
    fmt.keep_with_next = keep or None
    return paragraph


def rules(paragraph, *sides: tuple[str, str, str, str], fill: str | None = None) -> None:
    """Paragraph borders, each (side, size in eighths of a point, space, color), and a fill."""
    properties, edges = paragraph._p.get_or_add_pPr(), make("w:pBdr")
    for side, size, space, color in sides:
        edges.append(make(f"w:{side}", val="single", sz=size, space=space, color=color))
    properties.insert_element_before(edges, *AFTER_PARA_BORDER)
    if fill:
        properties.insert_element_before(make("w:shd", val="clear", color="auto", fill=fill),
                                         *AFTER_PARA_BORDER[1:])


def add_field(paragraph, instruction: str, look: dict) -> None:
    """A page-number field: begin, instruction, separate, placeholder, end."""
    for tag, value in (("w:fldChar", "begin"), ("w:instrText", f" {instruction} "),
                       ("w:fldChar", "separate"), ("w:t", "1"), ("w:fldChar", "end")):
        run = paragraph.add_run()
        style(run, **look)
        if tag == "w:fldChar":
            run._element.append(make(tag, fldCharType=value))
            continue
        node = OxmlElement(tag)
        node.set(qn("xml:space"), "preserve")
        node.text = value
        run._element.append(node)


# ---------------------------------------------------------------- tables
def set_table(table, widths: list[int], borders: bool, rule: str | None = None) -> None:
    """Fixed column widths; hairline borders, or none with an optional accent rule underneath."""
    table.autofit = False
    properties = table._tbl.tblPr
    properties.find(qn("w:tblW")).attrib.update({qn("w:w"): str(sum(widths)), qn("w:type"): "dxa"})
    edges = make("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if borders or (rule and edge == "bottom"):
            edges.append(make(f"w:{edge}", val="single", sz="4" if borders else "12", space="0",
                              color=HAIRLINE if borders else rule))
        else:
            edges.append(make(f"w:{edge}", val="nil"))
    properties.insert_element_before(edges, *AFTER_TABLE_BORDERS)
    margins = make("w:tblCellMar")
    pads = (50, 115, 50, 115) if borders else (0, 0, 110 if rule else 0, 0)
    for edge, value in zip(("top", "left", "bottom", "right"), pads):
        margins.append(make(f"w:{edge}", w=str(value), type="dxa"))
    properties.insert_element_before(margins, "w:tblLook")
    for column, width in zip(table.columns, widths):
        column.width = Emu(width * 635)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Emu(width * 635)


def keep_row(row, header=False) -> None:
    """A row never splits across pages; the header row repeats on each page."""
    properties = row._tr.get_or_add_trPr()
    properties.append(make("w:cantSplit"))
    if header:
        properties.append(make("w:tblHeader"))


def shade(cell, fill: str) -> None:
    node = make("w:shd", val="clear", color="auto", fill=fill)
    cell._tc.get_or_add_tcPr().insert_element_before(node, *AFTER_CELL_SHADE)


def column_widths(rows: list[list[str]], total: int) -> list[int]:
    """Twips: shares for the table's pattern, then no column narrower than its longest word."""
    words = [[plain(cell) for cell in row] for row in rows]
    count = len(words[0])
    if count == 5 and words[0][0].strip() == "#":
        shares = list(ACTION_WIDTHS)
    else:
        longest = [max(len(row[i]) for row in words) + 4 for i in range(count)]
        shares = [min(max(size / sum(longest), 0.10), 0.50) for size in longest]
    widths = [int(total * share / sum(shares)) for share in shares]
    widths[-1] += total - sum(widths)
    # About 0.6 em a character at 9.5 pt (114 twips), capitals at 8.5 pt tracked (132), plus padding.
    need = [max(max((len(word) for row in words[1:] for word in row[i].split()), default=0) * 114,
                len(words[0][i]) * 132) + 260 for i in range(count)]
    for i in range(count):
        short = need[i] - widths[i]
        donor = max(range(count), key=lambda j: widths[j] - need[j])
        if short > 0 and donor != i and widths[donor] - need[donor] >= short:
            widths[donor] -= short
            widths[i] += short
    return widths


def add_table(document, rows: list[list[str]], width: int, colors: tuple[str, str], look: dict):
    fill, on_fill = colors
    owner = [i for i, cell in enumerate(rows[0]) if cell.strip().lower() == "owner"]
    table = document.add_table(rows=0, cols=len(rows[0]))
    for position, cells in enumerate(rows):
        row = table.add_row()
        keep_row(row, header=position == 0)
        for column, (cell, text) in enumerate(zip(row.cells, cells)):
            paragraph = cell.paragraphs[0]
            if position == 0:
                shade(cell, fill)
                style(paragraph.add_run(plain(text).upper()), look["font"], 8.5, on_fill,
                      bold=True, spacing=10)
            else:
                add_text(paragraph, text, look, bold=column in owner)
            spacing(paragraph, after=0, line=1.1)
    set_table(table, column_widths(rows, width), borders=True)
    spacing(document.add_paragraph(), after=6, line=Pt(1))


# ---------------------------------------------------------------- the page
def build(title_text: str, facts: list[tuple[str, str]], sections: list, brand: dict):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = PAGE_SIZES[brand["page_size"].lower()]
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(0.8)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.4)
    width = int((section.page_width - section.left_margin - section.right_margin) / 635)

    # The accent fills; text in it is darkened in its own hue until it reads, 4.5:1 on white
    # (3:1 for bullets and rules), so even a pale brand color gives legible headings.
    accent, ink = brand["accent_color"], readable(brand["text_color"])
    strong, mark, tint = readable(accent), readable(accent, ratio=3), pale(accent)
    on_fill = "FFFFFF" if contrast("FFFFFF", accent) >= 4.5 else readable(accent, accent)
    display, body = brand["display_font"], brand["body_font"]
    text_look = {"font": body, "size": 10.5, "color": ink, "link": strong}

    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = body, Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(ink)
    properties = document.core_properties
    properties.title, properties.author = title_text, brand["organization"]
    properties.last_modified_by = brand["organization"]
    properties.comments = properties.subject = properties.keywords = properties.category = ""

    # Letterhead: logo, organization and tagline at the left, the document label at the right.
    band = document.add_table(rows=1, cols=2)
    left, right = band.rows[0].cells
    stack = []
    if brand["logo_path"]:
        picture = Image.from_file(str(brand["logo_path"]))
        wide = LOGO_HEIGHT * picture.px_width / picture.px_height > LOGO_MAX_WIDTH
        size = {"width": LOGO_MAX_WIDTH} if wide else {"height": LOGO_HEIGHT}
        left.paragraphs[0].add_run().add_picture(str(brand["logo_path"]), **size)
        stack.append(spacing(left.paragraphs[0], after=4, line=1.0))
    for key, look in (("organization", dict(font=display, size=13, color=ink, bold=True)),
                      ("tagline", dict(font=display, size=9.5, color=strong, italic=True))):
        if brand[key]:
            paragraph = left.add_paragraph() if stack else left.paragraphs[0]
            style(paragraph.add_run(brand[key]), **look)
            stack.append(spacing(paragraph, after=1, line=1.0))
    label = right.paragraphs[0]
    style(label.add_run(brand["document_label"]), display, 15, strong, bold=True, spacing=30)
    spacing(label, after=0, line=1.0).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    set_table(band, [int(width * 0.6), width - int(width * 0.6)], borders=False, rule=mark)
    spacing(document.add_paragraph(), after=12, line=Pt(1))

    title = document.add_paragraph()
    style(title.add_run(title_text), display, 18, strong, bold=True)
    spacing(title, after=10, line=1.05)

    if facts:
        table = document.add_table(rows=0, cols=2)
        for key, value in facts:
            row = table.add_row()
            keep_row(row)
            name, content = row.cells
            shade(name, tint)
            style(name.paragraphs[0].add_run(key.upper()), body, 8.5, GRAY, bold=True, spacing=10)
            spacing(name.paragraphs[0], after=0, line=1.1)
            add_text(content.paragraphs[0], value, {**text_look, "size": 10})
            spacing(content.paragraphs[0], after=0, line=1.1)
        set_table(table, [int(width * 0.22), width - int(width * 0.22)], borders=True)
        spacing(document.add_paragraph(), after=8, line=Pt(1))

    for heading, blocks in sections:
        paragraph = document.add_paragraph()
        style(paragraph.add_run(heading), display, 12.5, strong, bold=True)
        spacing(paragraph, before=14, after=5, keep=True)
        note = heading.lower() == AI_NOTE.lower()  # the AI statement, set apart as a block
        for block in blocks:
            if block.kind == "table":
                add_table(document, block.rows, width, (accent, on_fill), {**text_look, "size": 9.5})
                continue
            paragraph = document.add_paragraph()
            if block.kind == "para":
                add_text(paragraph, block.text, {**text_look, "size": 9.5} if note else text_look)
                spacing(paragraph)
                if note:
                    rules(paragraph, ("top", "4", "5", tint), ("left", "24", "9", mark),
                          ("bottom", "4", "5", tint), ("right", "4", "9", tint), fill=tint)
                    fmt = paragraph.paragraph_format
                    fmt.left_indent, fmt.right_indent = Pt(11), Pt(10)
                continue
            number = block.kind == "number"
            symbol, color = (f"{block.number}.", strong) if number else ("•", mark)
            indent = Inches(0.3 if number else 0.25)
            style(paragraph.add_run(symbol + "\t"), body, 10.5, color, bold=number)
            add_text(paragraph, block.text, text_look)
            fmt = paragraph.paragraph_format
            fmt.left_indent, fmt.first_line_indent = indent, -indent
            fmt.tab_stops.add_tab_stop(indent)
            spacing(paragraph, after=3)

    # Footer: the brand line at the left, the page count at the right. Normal style, because the
    # built-in Footer style carries a center tab that would pull the page count to the middle.
    footer = section.footer.paragraphs[0]
    footer.style = document.styles["Normal"]
    fine = {"font": body, "size": 8, "color": GRAY}
    rules(footer, ("top", "4", "6", mark))
    footer.paragraph_format.tab_stops.add_tab_stop(Emu(width * 635), WD_TAB_ALIGNMENT.RIGHT)
    style(footer.add_run(brand["footer_text"] + "\tPage "), **fine)
    add_field(footer, "PAGE", fine)
    style(footer.add_run(" of "), **fine)
    add_field(footer, "NUMPAGES", fine)
    spacing(footer, after=0, line=1.0)
    return document


# ---------------------------------------------------------------- PDF
def find_soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    return found or next((path for path in KNOWN_SOFFICE if Path(path).is_file()), None)


def convert(docx: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise Fail(f"LibreOffice (soffice) was not found. {docx.name} is left beside the recap: "
                   "convert this file to PDF with your document tools, then read the pages.")
    pdf = docx.with_suffix(".pdf")
    pdf.unlink(missing_ok=True)  # an old render must never pass for this one
    retry = f"Close any open LibreOffice windows and run again; {docx.name} is left beside the recap."
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
            capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired:
        raise Fail(f"LibreOffice did not finish within three minutes. {retry}") from None
    if result.returncode != 0 or not pdf.is_file():
        raise Fail(f"LibreOffice did not produce {pdf.name}. {retry}")
    return pdf


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Build the PDF of a meeting recap.")
    parser.add_argument("recap", type=Path, help="the markdown recap")
    parser.add_argument("--brand", type=Path, default=SKILL_DIR / "brand",
                        help="folder holding brand.md and the logo (default: the skill's brand/)")
    args = parser.parse_args(argv)
    recap_path = args.recap.expanduser().resolve()
    if not recap_path.is_file():
        raise Fail(f"no recap at {recap_path}")
    brand = load_brand(args.brand.expanduser().resolve())
    title, facts, sections = parse(recap_path)
    docx = recap_path.with_suffix(".docx")
    build(title, facts, sections, brand).save(str(docx))
    pdf = convert(docx)
    docx.unlink()
    print(f"Wrote {pdf}")
    print("Next: render every page to an image and look at each one before sending.")


if __name__ == "__main__":
    try:
        main()
    except Fail as problem:
        sys.exit(f"build_recap: {problem}")
